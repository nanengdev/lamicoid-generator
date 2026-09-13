import streamlit as st
import pandas as pd
import math
import zipfile
import io
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

st.set_page_config(page_title='Generador Lamicoides', layout='wide')
st.title('🛠️ Generador de Etiquetas Lamicoides')

# --- CONFIGURACIÓN EN SIDEBAR ---
st.sidebar.header('Parámetros de Diseño')
dpi = st.sidebar.number_input('DPI (LaserGRBL)', value=600)
sep_mm = st.sidebar.slider('Separación (mm)', 0.0, 10.0, 3.0)
margen_hoja = st.sidebar.slider('Margen Hoja (mm)', 0.0, 20.0, 5.0)
cols_num = st.sidebar.number_input('Columnas', min_value=1, value=3)
margen_seg = st.sidebar.slider('Margen Seguridad Texto (mm)', 0.0, 5.0, 1.5)

def mm_a_px(mm): return int(round(mm * dpi / 25.4))

def ajustar_fuente(draw, lineas, w_px, h_px):
    for pt in range(20, 6, -1):
        fnt = ImageFont.load_default()
        max_w = 0
        for l in lineas:
            b = draw.textbbox((0,0), l, font=fnt)
            max_w = max(max_w, b[2]-b[0])
        if max_w <= w_px: return fnt
    return ImageFont.load_default()

def distribuir(etiquetas):
    cols = min(cols_num, len(etiquetas))
    filas = math.ceil(len(etiquetas) / cols)
    anchos_cols = [0.0] * cols
    altos_filas = [0.0] * filas
    for i, et in enumerate(etiquetas):
        c, f = i % cols, i // cols
        anchos_cols[c] = max(anchos_cols[c], et['ancho_mm'])
        altos_filas[f] = max(altos_filas[f], et['alto_mm'])
    px, py = [], []
    cx, cy = margen_hoja, margen_hoja
    for a in anchos_cols: px.append(cx); cx += a + sep_mm
    for a in altos_filas: py.append(cy); cy += a + sep_mm
    for i, et in enumerate(etiquetas):
        c, f = i % cols, i // cols
        et['x_mm'] = px[c] + (anchos_cols[c] - et['ancho_mm'])/2
        et['y_mm'] = py[f] + (altos_filas[f] - et['alto_mm'])/2
    return cx - sep_mm + margen_hoja, cy - sep_mm + margen_hoja

uploaded = st.file_uploader('Carga tu archivo Excel (.xlsx)', type=['xlsx'])

if uploaded:
    wb = load_workbook(uploaded, data_only=True)
    sh_name = st.selectbox('Selecciona la Hoja', wb.sheetnames)
    
    if st.button('🚀 Generar Archivos para Láser'):
        ws = wb[sh_name]
        header = {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value}
        
        etiquetas = []
        for r in range(2, ws.max_row + 1):
            t1 = str(ws.cell(r, header.get('Texto1', 1)).value or "").strip()
            t2 = str(ws.cell(r, header.get('Texto2', 2)).value or "").strip()
            t3 = str(ws.cell(r, header.get('Texto3', 3)).value or "").strip()
            if not any([t1, t2, t3]): continue
            
            lineas = [l for l in [t1, t2, t3] if l]
            ancho = float(ws.cell(r, header.get('Ancho_mm', 4)).value or 50)
            alto = float(ws.cell(r, header.get('Alto_mm', 5)).value or 20)
            cant = int(ws.cell(r, header.get('Cantidad', 6)).value or 1)
            
            for _ in range(cant):
                etiquetas.append({'lineas': lineas, 'ancho_mm': ancho, 'alto_mm': alto})

        if etiquetas:
            w_h, h_h = distribuir(etiquetas)
            img = Image.new("L", (mm_a_px(w_h), mm_a_px(h_h)), 255)
            draw = ImageDraw.Draw(img)
            svg = [f'<?xml version="1.0"?><svg width="{w_h}mm" height="{h_h}mm" viewBox="0 0 {w_h} {h_h}" xmlns="http://www.w3.org/2000/svg">']
            svg.append(f'<rect x="0" y="0" width="{w_h}" height="{h_h}" fill="none" stroke="none"/>')
            for et in etiquetas:
                x_px, y_px = mm_a_px(et['x_mm']), mm_a_px(et['y_mm'])
                w_px, h_px = mm_a_px(et['ancho_mm']), mm_a_px(et['alto_mm'])
                m_px = mm_a_px(margen_seg)
                fnt = ajustar_fuente(draw, et['lineas'], w_px - 2*m_px, h_px - 2*m_px)
                y_offset = y_px + m_px
                for ln in et['lineas']:
                    draw.text((x_px + m_px, y_offset), ln, font=fnt, fill=0)
                    y_offset += mm_a_px(5)
                svg.append(f'<rect x="{et["x_mm"]}" y="{et["y_mm"]}" width="{et["ancho_mm"]}" height="{et["alto_mm"]}" fill="none" stroke="red" stroke-width="0.1"/>')
            svg.append('</svg>')
            png_out = io.BytesIO()
            img.save(png_out, format='PNG', dpi=(dpi, dpi))
            zip_buf = io.BytesIO()
            with zipfile.ZipFile(zip_buf, 'w') as zf:
                zf.writestr('01_grabado.png', png_out.getvalue())
                zf.writestr('02_corte.svg', "\\n".join(svg))
            st.success(f'¡Éxito! Generadas {len(etiquetas)} etiquetas.')
            st.download_button('🎁 Descargar ZIP para LaserGRBL', zip_buf.getvalue(), 'lamicoides_laser.zip')
